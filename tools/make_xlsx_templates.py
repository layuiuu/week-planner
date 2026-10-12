"""生成 samples/ 下的 Excel 模板(需要用 openpyxl)。

用法:
  pip install openpyxl
  python tools/make_xlsx_templates.py

输出:
  samples/week-template.xlsx   周表模板
  samples/tasks-template.xlsx  任务模板

登录页/设置页里的「下载 Excel 模板」按钮直接指向这两个文件。
(仓库里已经生成好了,不改模板内容的话不用重新跑这个脚本。)
"""

import datetime as dt
import os
import sys

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("需要先安装 openpyxl:pip install openpyxl")

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "samples")

HEAD_FILL = PatternFill("solid", fgColor="4F46E5")
HEAD_FONT = Font(color="FFFFFF", bold=True)


def style_sheet(ws, headers, widths):
    for col, (title, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.freeze_panes = "A2"


def monday_of_this_week():
    today = dt.date.today()
    return today - dt.timedelta(days=today.weekday())


def make_week_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "周表"
    headers = ["星期", "板块名称", "类别", "开始", "结束"]
    style_sheet(ws, headers, [10, 16, 12, 10, 10])

    rows = [
        ("周一", "上午学习", "学习", dt.time(9, 0), dt.time(12, 0)),
        ("周一", "下午工作", "工作", dt.time(14, 0), dt.time(16, 0)),
        ("周一", "力量训练", "健身", dt.time(19, 0), dt.time(20, 0)),
        ("周二", "上午学习", "学习", dt.time(9, 0), dt.time(12, 0)),
        ("周二", "下午工作", "工作", dt.time(14, 0), dt.time(18, 0)),
        ("周三", "上午学习", "学习", dt.time(9, 0), dt.time(12, 0)),
        ("周三", "夜跑", "健身", dt.time(19, 0), dt.time(20, 30)),
        ("周四", "上午学习", "学习", dt.time(9, 0), dt.time(12, 0)),
        ("周四", "下午工作", "工作", dt.time(14, 0), dt.time(16, 0)),
        ("周五", "上午学习", "学习", dt.time(9, 0), dt.time(12, 0)),
        ("周五", "下午工作", "工作", dt.time(14, 0), dt.time(16, 0)),
        ("周五", "睡前阅读", "学习", dt.time(20, 0), dt.time(21, 0)),
        ("周六", "大扫除", "生活", dt.time(10, 0), dt.time(12, 0)),
        ("周六", "兴趣项目", "兴趣", dt.time(15, 0), dt.time(17, 0)),
        ("周日", "周复盘", "生活", dt.time(10, 0), dt.time(11, 0)),
        ("周日", "力量训练", "健身", dt.time(15, 0), dt.time(16, 30)),
    ]
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c >= 4:
                cell.number_format = "HH:MM"
                cell.alignment = Alignment(horizontal="center")
    path = os.path.join(OUT, "week-template.xlsx")
    wb.save(path)
    print("wrote", path)


def make_tasks_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "任务"
    headers = ["日期", "板块名称", "任务", "备注"]
    style_sheet(ws, headers, [14, 16, 30, 20])

    monday = monday_of_this_week()
    rows = [
        (monday, "上午学习", "复习第三章 3.1–3.3", "重点看例题"),
        (monday, "上午学习", "整理上周错题", ""),
        (monday, "下午工作", "写周报", "周五交"),
        (monday + dt.timedelta(days=1), "上午学习", "做完第五章练习", ""),
        (monday + dt.timedelta(days=2), "上午学习", "背 50 个单词", ""),
        (monday + dt.timedelta(days=2), "夜跑", "慢跑 5 公里", ""),
        (monday + dt.timedelta(days=4), "睡前阅读", "读 20 分钟", ""),
        (monday + dt.timedelta(days=5), "大扫除", "擦地板 + 换床单", ""),
        (monday + dt.timedelta(days=6), "周复盘", "写这周的复盘", "用「复盘」页"),
    ]
    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            cell = ws.cell(row=r, column=c, value=value)
            if c == 1:
                cell.number_format = "yyyy-mm-dd"
                cell.alignment = Alignment(horizontal="center")
    path = os.path.join(OUT, "tasks-template.xlsx")
    wb.save(path)
    print("wrote", path)


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    make_week_template()
    make_tasks_template()
